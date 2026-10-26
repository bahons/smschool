import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def build_presentation(output_pptx: str = "outputs/presentation_assets/SmartSchool_Demo_Presentation.pptx"):
    """
    Builds a 6-slide executive presentation using python-pptx:
    1. Мәселе (Problem)
    2. Шешім (Solution)
    3. Жүйе архитектурасы (Architecture & Pose AI)
    4. Демо нәтижелері (Demo Results)
    5. Салыстырмалы анализ (Comparison)
    6. Қорытынды және болашағы (Conclusion & Roadmap)
    """
    os.makedirs(os.path.dirname(output_pptx), exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_header(slide, title_text, subtitle_text=""):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(26, 35, 126)  # Deep Navy

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(14)
            p2.font.color.rgb = RGBColor(100, 116, 139)

    # ---------------- SLIDE 1: Титул және Мәселе ----------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_header(slide1, "Smart School: Компьютерлік көруге негізделген қауіпсіздік жүйесі", "Мектептердегі буллинг пен төбелесті ерте анықтау платформасы")
    tb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🔴 Өзекті мәселелер:"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(198, 40, 40)

    bullets1 = [
        "Мектеп коридорлары мен сыныптарындағы конфликттер мен төбелестер дер кезінде байқалмайды.",
        "Кезекші мұғалімдер мен күзет барлық 10-50 камераны бір уақытта адам факторына байланысты бақылай алмайды.",
        "Дәстүрлі жүйелер тек жазба сақтайды (пост-фактум), ал алдын-алу немесе оқиға кезінде жедел хабарлау жоқ.",
        "Жеке құпиялылық (Privacy): Оқушылардың бет-бейнесін сақтамай, этикалық тұрғыдан қауіпсіз болу қажеттілігі."
    ]
    for b in bullets1:
        p_b = tf1.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(18)
        p_b.font.color.rgb = RGBColor(40, 40, 40)

    # ---------------- SLIDE 2: Шешім ----------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_header(slide2, "Біздің шешім: Нақты уақыттағы биомеханикалық талдау", "Бетті танымай, тек қимыл-қозғалыс пен кинетиканы бағалау")
    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    bullets2 = [
        "🛡 Жоғары сенімділік: Тек силуэт пен қаңқа сызықтары (17 COCO keypoints) негізінде сараптама.",
        "🎯 Conflict Index метрикасы: Білектер (wrists) жылдамдығы мен үдеуі + оқушылар арасындағы арақашықтық.",
        "⏱ 2 секундтық растау сүзгісі: Жалған дабылдарды (қол көтеріп амандасу, жылдам жүру) болдырмау үшін индекс 2 сек бойы қауіп шегінен асуы шарт.",
        "🔒 Құпиялылық (Privacy First): Бас аймағы бұлдырланады, бетті тану қолданылмайды."
    ]
    for b in bullets2:
        p_b = tf2.add_paragraph()
        p_b.text = "✔ " + b
        p_b.font.size = Pt(18)
        p_b.font.color.rgb = RGBColor(27, 94, 32)

    # ---------------- SLIDE 3: Архитектура ----------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_header(slide3, "Жүйе архитектурасы (Edge AI & Pipeline)", "Камера ағынынан бастап Telegram/Dashboard хабарламасына дейін")
    tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.5))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    arch_steps = [
        "1. Video Stream Ingestion (RTSP / MP4 File Stream) → нақты FPS және уақыт белгісін (timestamp) алу.",
        "2. YOLOv8n-pose Detector → адамдардың 17 негізгі буындарын (keypoints) анықтау.",
        "3. ByteTrack Tracker → әр адамға кадрлар бойы тұрақты ID тағайындау.",
        "4. Conflict Index Engine → білектер үдеуі + жақындық + KDE тығыздық динамикасы.",
        "5. State Machine → НОРМА (жасыл) → ЕСКЕРТУ (сары) → ҚАУІП: ТӨБЕЛЕС (қызыл).",
        "6. Alert & Notification Module → Telegram Bot жедел хабарламасы + Web Dashboard."
    ]
    for s in arch_steps:
        p_s = tf3.add_paragraph()
        p_s.text = s
        p_s.font.size = Pt(16)
        p_s.font.color.rgb = RGBColor(30, 41, 59)

    # ---------------- SLIDE 4: Демо нәтижелері ----------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_header(slide4, "Демо және Тестілеу нәтижелері (outputs/)", "Тыныш сынып және Нақты төбелес бейнелері бойынша")

    # Add images if available
    img_calm = "outputs/video2_calm/conflict_timeline.png"
    img_fight = "outputs/video1_fight/conflict_timeline.png"

    if os.path.exists(img_calm):
        slide4.shapes.add_picture(img_calm, Inches(0.8), Inches(2.0), width=Inches(5.6))
    if os.path.exists(img_fight):
        slide4.shapes.add_picture(img_fight, Inches(6.8), Inches(2.0), width=Inches(5.6))

    # ---------------- SLIDE 5: Салыстырмалы Анализ ----------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_header(slide5, "Салыстырмалы Анализ (Инфографика)", "Екі жағдайдың параметрлік және уақыттық айырмашылығы")
    comp_img = "outputs/comparison.png"
    if os.path.exists(comp_img):
        slide5.shapes.add_picture(comp_img, Inches(1.2), Inches(1.8), width=Inches(10.8))

    # ---------------- SLIDE 6: Қорытынды ----------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_header(slide6, "Қорытынды және Болашақ жоспар", "Жүйені мектептерге енгізу тиімділігі")
    tb6 = slide6.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.5))
    tf6 = tb6.text_frame
    tf6.word_wrap = True
    conclusions = [
        "🚀 Жоғары жылдамдық: YOLOv8n жеңіл моделі кез келген Edge серверде (30+ FPS) еркін жұмыс істейді.",
        "🎯 Нөлдік жалған дабыл: 2 секундтық қауіп фильтрі тыныш сыныпта бірде-бір жалған дабыл тудырмады.",
        "📲 Мұғалімдер мен күзетке 1-2 секунд ішінде фото-дәлелмен нақты хабарлама жетеді.",
        "📈 Келесі қадам: Мектептердің орталықтандырылған 102/SOS желісімен интеграциялау."
    ]
    for c in conclusions:
        p_c = tf6.add_paragraph()
        p_c.text = "★ " + c
        p_c.font.size = Pt(18)
        p_c.font.color.rgb = RGBColor(13, 71, 161)

    prs.save(output_pptx)
    print(f"[SUCCESS] 6 слайдтық презентация сақталды: {output_pptx}")
    return output_pptx
